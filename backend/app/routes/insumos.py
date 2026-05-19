from fastapi import APIRouter, Depends, Header, HTTPException, Request, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from pydantic import BaseModel
from typing import Optional
from datetime import datetime, timedelta, timezone
import pyotp
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.models.insumo import Insumo, TipoInsumo
from app.models.movimiento import Movimiento, TipoMovimiento
from app.models.usuario import Usuario, RolUsuario
from app.schemas.insumo import InsumoCreate, InsumoUpdate, InsumoResponse
from app.schemas.comun import PaginatedResponse
from app.utils.deps import get_usuario_actual, require_operador, require_admin
from app.utils.auditoria import registrar, get_ip

router = APIRouter(prefix="/insumos", tags=["Insumos"])

TOTP_VALID_WINDOW = 1


class InsumoAlerta(BaseModel):
    id: int
    nombre: str
    stock_actual: int
    stock_minimo: int
    deficit: int
    sala: str | None
    categoria: str | None
    tipo: str = "insumo"

    class Config:
        from_attributes = True


def _build_query(db: Session,
                 nombre: Optional[str] = None,
                 sala_id: Optional[int] = None,
                 categoria_id: Optional[int] = None,
                 bajo_stock: Optional[bool] = None,
                 incluir_inactivos: bool = False,
                 tipo: Optional[TipoInsumo] = None):
    """Construye la query base con filtros opcionales reutilizable.
    Por defecto excluye insumos desactivados (soft-delete).
    """
    q = db.query(Insumo)
    if not incluir_inactivos:
        q = q.filter(Insumo.activo.is_(True))
    if nombre:
        q = q.filter(Insumo.nombre.ilike(f"%{nombre}%"))
    if sala_id is not None:
        q = q.filter(Insumo.sala_id == sala_id)
    if categoria_id is not None:
        q = q.filter(Insumo.categoria_id == categoria_id)
    if bajo_stock:
        q = q.filter(Insumo.stock_actual <= Insumo.stock_minimo)
    if tipo is not None:
        q = q.filter(Insumo.tipo == tipo)
    return q


# ---------------------------------------------------------------------------
# IMPORTANTE: rutas estaticas deben ir ANTES de la ruta dinamica /{insumo_id},
# o FastAPI las interpreta como un entero y devuelve 422.
# ---------------------------------------------------------------------------

@router.get("/alertas", response_model=list[InsumoAlerta])
def alertas_stock(
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual)
):
    """Insumos activos cuyo stock_actual es menor o igual al stock_minimo."""
    insumos = (
        db.query(Insumo)
        .filter(Insumo.activo.is_(True))
        .filter(Insumo.stock_actual <= Insumo.stock_minimo)
        .order_by(Insumo.stock_actual.asc())
        .all()
    )
    return [
        InsumoAlerta(
            id=i.id,
            nombre=i.nombre,
            stock_actual=i.stock_actual,
            stock_minimo=i.stock_minimo,
            deficit=i.stock_minimo - i.stock_actual,
            sala=i.sala.nombre if i.sala else None,
            categoria=i.categoria.nombre if i.categoria else None,
            tipo=i.tipo.value if i.tipo else "insumo",
        )
        for i in insumos
    ]


@router.get("/alertas-resueltas", response_model=list[InsumoAlerta])
def alertas_resueltas(
    dias: int = 30,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual)
):
    """Insumos activos que YA superaron el stock minimo y ademas tuvieron al
    menos una entrada en los ultimos `dias` dias."""
    desde = datetime.now(timezone.utc) - timedelta(days=dias)

    subq = (
        db.query(Movimiento.insumo_id)
        .filter(
            Movimiento.tipo == TipoMovimiento.entrada,
            Movimiento.fecha >= desde,
        )
        .distinct()
        .subquery()
    )

    insumos = (
        db.query(Insumo)
        .filter(
            Insumo.activo.is_(True),
            Insumo.stock_actual > Insumo.stock_minimo,
            Insumo.id.in_(subq),
        )
        .order_by(Insumo.nombre)
        .all()
    )

    return [
        InsumoAlerta(
            id=i.id,
            nombre=i.nombre,
            stock_actual=i.stock_actual,
            stock_minimo=i.stock_minimo,
            deficit=i.stock_minimo - i.stock_actual,
            sala=i.sala.nombre if i.sala else None,
            categoria=i.categoria.nombre if i.categoria else None,
            tipo=i.tipo.value if i.tipo else "insumo",
        )
        for i in insumos
    ]


@router.get("/exportar")
def exportar_insumos(
    formato: str = "csv",
    nombre: Optional[str] = None,
    sala_id: Optional[int] = None,
    categoria_id: Optional[int] = None,
    bajo_stock: Optional[bool] = None,
    incluir_inactivos: bool = False,
    tipo: Optional[TipoInsumo] = None,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual)
):
    """Exporta el inventario como CSV o Excel con los filtros activos.

    formato=csv  -> archivo .csv con BOM UTF-8 (abre bien en Excel).
    formato=xlsx -> archivo .xlsx con cabecera coloreada y columnas ajustadas.
    """
    insumos = _build_query(
        db, nombre, sala_id, categoria_id, bajo_stock, incluir_inactivos, tipo
    ).order_by(Insumo.nombre).all()

    cabecera = [
        "Nombre", "Tipo", "SKU", "Codigo de barras", "Descripcion",
        "Stock actual", "Stock minimo", "Costo unitario",
        "Sala", "Categoria", "Estado"
    ]

    def _estado(i: Insumo) -> str:
        if not i.activo:
            return "inactivo"
        if i.stock_actual == 0:
            return "agotado"
        if i.stock_actual <= i.stock_minimo:
            return "bajo_stock"
        return "ok"

    filas = [
        [
            i.nombre,
            i.tipo.value if i.tipo else "insumo",
            i.sku or "",
            i.codigo_barras or "",
            i.descripcion or "",
            i.stock_actual,
            i.stock_minimo,
            float(i.costo_unitario) if i.costo_unitario is not None else "",
            i.sala.nombre if i.sala else "",
            i.categoria.nombre if i.categoria else "",
            _estado(i),
        ]
        for i in insumos
    ]

    if formato == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Inventario"

        header_fill = PatternFill("solid", fgColor="0F766E")  # teal-700
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx, titulo in enumerate(cabecera, 1):
            cell = ws.cell(row=1, column=col_idx, value=titulo)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for fila in filas:
            ws.append(fila)

        for col_idx, _ in enumerate(cabecera, 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                (len(str(ws.cell(row=r, column=col_idx).value or ""))
                 for r in range(1, ws.max_row + 1)),
                default=10,
            )
            ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=inventario_hestia.xlsx"},
        )

    # CSV con BOM UTF-8 para compatibilidad con Excel en Windows
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(cabecera)
    writer.writerows(filas)
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=inventario_hestia.csv"}
    )


@router.get("/sugerencias", response_model=list[str])
def sugerencias_insumos(
    q: str,
    limit: int = 8,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual),
):
    """Devuelve nombres de insumos activos que coinciden con `q`.

    Usado por el componente SearchWithSuggestions del frontend para
    mostrar un dropdown de autocompletado mientras el usuario escribe.
    Requiere minimo 2 caracteres para evitar queries triviales.
    Solo devuelve nombres (strings), no el objeto completo.
    """
    if len(q.strip()) < 2:
        return []
    resultados = (
        db.query(Insumo.nombre)
        .filter(
            Insumo.activo.is_(True),
            Insumo.nombre.ilike(f"%{q.strip()}%"),
        )
        .order_by(Insumo.nombre)
        .limit(min(limit, 15))
        .all()
    )
    return [r[0] for r in resultados]


@router.get("/", response_model=PaginatedResponse[InsumoResponse])
def listar_insumos(
    skip: int = 0,
    limit: int = 20,
    nombre: Optional[str] = None,
    sala_id: Optional[int] = None,
    categoria_id: Optional[int] = None,
    bajo_stock: Optional[bool] = None,
    incluir_inactivos: bool = False,
    tipo: Optional[TipoInsumo] = None,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual)
):
    q = _build_query(db, nombre, sala_id, categoria_id, bajo_stock, incluir_inactivos, tipo)
    total = q.count()
    insumos = q.offset(skip).limit(limit).all()
    return {"total": total, "skip": skip, "limit": limit, "data": insumos}


@router.get("/{insumo_id}", response_model=InsumoResponse)
def obtener_insumo(
    insumo_id: int,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual)
):
    insumo = db.query(Insumo).filter(Insumo.id == insumo_id).first()
    if not insumo:
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    return insumo


@router.post("/", response_model=InsumoResponse)
def crear_insumo(
    request: Request,
    insumo: InsumoCreate,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_operador)
):
    nuevo = Insumo(**insumo.model_dump())
    db.add(nuevo)
    db.flush()  # asigna ID sin hacer commit; necesario para auto-generar SKU
    if not nuevo.sku:
        nuevo.sku = f"HST-{nuevo.id:05d}"
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        detalle = str(exc.orig).lower()
        if "sku" in detalle:
            raise HTTPException(
                status_code=409, detail="El SKU ya esta en uso por otro insumo."
            )
        if "codigo_barras" in detalle:
            raise HTTPException(
                status_code=409, detail="El codigo de barras ya esta en uso."
            )
        raise HTTPException(status_code=409, detail="Conflicto: dato duplicado.")
    db.refresh(nuevo)
    registrar(
        db, "CREAR_INSUMO", usuario=usuario,
        entidad="insumo", entidad_id=nuevo.id,
        detalle=nuevo.nombre, ip=get_ip(request),
    )
    return nuevo


@router.put("/{insumo_id}", response_model=InsumoResponse)
def actualizar_insumo(
    insumo_id: int,
    request: Request,
    datos: InsumoUpdate,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_operador)
):
    insumo = db.query(Insumo).filter(Insumo.id == insumo_id).first()
    if not insumo:
        raise HTTPException(status_code=404, detail="Insumo no encontrado")

    cambio_activo: str | None = None
    if datos.activo is not None and datos.activo != insumo.activo:
        if usuario.rol != RolUsuario.admin:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Solo administradores pueden cambiar el estado activo.",
            )
        cambio_activo = (
            "REACTIVAR_INSUMO" if datos.activo else "DESACTIVAR_INSUMO"
        )

    for campo, valor in datos.model_dump(exclude_unset=True).items():
        setattr(insumo, campo, valor)

    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        detalle = str(exc.orig).lower()
        if "sku" in detalle:
            raise HTTPException(
                status_code=409, detail="El SKU ya esta en uso por otro insumo."
            )
        if "codigo_barras" in detalle:
            raise HTTPException(
                status_code=409, detail="El codigo de barras ya esta en uso."
            )
        raise HTTPException(status_code=409, detail="Conflicto: dato duplicado.")

    db.refresh(insumo)
    registrar(
        db, "EDITAR_INSUMO", usuario=usuario,
        entidad="insumo", entidad_id=insumo.id,
        detalle=insumo.nombre, ip=get_ip(request),
    )
    if cambio_activo:
        registrar(
            db, cambio_activo, usuario=usuario,
            entidad="insumo", entidad_id=insumo.id,
            detalle=insumo.nombre, ip=get_ip(request),
        )
    return insumo


@router.delete("/{insumo_id}")
def eliminar_insumo(
    insumo_id: int,
    request: Request,
    codigo_totp: str = Header(alias="x-totp-code"),
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_admin)
):
    """Soft-delete: marca el insumo como inactivo (activo=false)."""
    if not usuario.totp_habilitado or not usuario.totp_secret:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Debes tener el 2FA activado para desactivar insumos."
        )
    totp = pyotp.TOTP(usuario.totp_secret)
    if not totp.verify(codigo_totp, valid_window=TOTP_VALID_WINDOW):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Codigo 2FA incorrecto. El insumo no fue desactivado."
        )
    insumo = db.query(Insumo).filter(Insumo.id == insumo_id).first()
    if not insumo:
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    if not insumo.activo:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=f"'{insumo.nombre}' ya esta inactivo",
        )
    insumo.activo = False
    db.commit()
    registrar(
        db, "DESACTIVAR_INSUMO", usuario=usuario,
        entidad="insumo", entidad_id=insumo.id,
        detalle=insumo.nombre, ip=get_ip(request),
    )
    return {"mensaje": f"Insumo '{insumo.nombre}' desactivado"}
