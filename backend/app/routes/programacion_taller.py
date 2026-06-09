from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import and_
from datetime import date
from typing import Optional
import io
import re
import datetime as dt

from app.database import get_db
from app.models.programacion_taller import ProgramacionTaller
from app.models.taller import Taller
from app.models.sala import Sala
from app.models.usuario import Usuario
from app.schemas.programacion_taller import (
    ProgramacionTallerCreate,
    ProgramacionTallerUpdate,
    ProgramacionTallerResponse,
    ImportarProgramacionResponse,
)
from app.utils.deps import get_usuario_actual, require_operador

router = APIRouter(prefix="/programacion", tags=["Programacion"])

# ---------------------------------------------------------------------------
# Tabla de meses en espanol para parseo de fechas texto
# ---------------------------------------------------------------------------
_MESES_ES = {
    'ene': 1, 'enero': 1,
    'feb': 2, 'febrero': 2,
    'mar': 3, 'marzo': 3,
    'abr': 4, 'abril': 4,
    'may': 5, 'mayo': 5,
    'jun': 6, 'junio': 6,
    'jul': 7, 'julio': 7,
    'ago': 8, 'agosto': 8,
    'sep': 9, 'sept': 9, 'septiembre': 9,
    'oct': 10, 'octubre': 10,
    'nov': 11, 'noviembre': 11,
    'dic': 12, 'diciembre': 12,
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _load(prog_id: int, db: Session) -> ProgramacionTaller:
    p = (
        db.query(ProgramacionTaller)
        .options(
            joinedload(ProgramacionTaller.taller).joinedload(
                Taller.asignatura
            ),
            joinedload(ProgramacionTaller.sala),
        )
        .filter(ProgramacionTaller.id == prog_id)
        .first()
    )
    if not p:
        raise HTTPException(status_code=404, detail="Programacion no encontrada")
    return p


def _to_response(p: ProgramacionTaller) -> ProgramacionTallerResponse:
    taller = p.taller
    asig = taller.asignatura if taller else None
    return ProgramacionTallerResponse(
        id=p.id,
        taller_id=p.taller_id,
        taller_nombre=taller.nombre if taller else None,
        asignatura_id=asig.id if asig else None,
        asignatura_nombre=asig.nombre if asig else None,
        sala_id=p.sala_id,
        sala_nombre=p.sala.nombre if p.sala else None,
        fecha=p.fecha,
        hora_inicio=p.hora_inicio,
        hora_fin=p.hora_fin,
        docente_nombre=p.docente_nombre,
        seccion=p.seccion,
        semestre=p.semestre,
        notas=p.notas,
        activo=p.activo,
    )


def _normalizar_hora(raw: object) -> Optional[str]:
    """Convierte multiples formatos de hora a 'HH:MM'.

    Soporta: '17:30:00', '17.30 A 18.50', '11:30-12:50', '8:31-9:50',
    objetos time de openpyxl (datetime.time), y floats Excel (0.729...).
    Devuelve None si no puede parsear.
    """
    if raw is None:
        return None
    if isinstance(raw, dt.time):
        return raw.strftime("%H:%M")
    if isinstance(raw, float):
        total_min = round(raw * 24 * 60)
        h, m = divmod(total_min, 60)
        return f"{h:02d}:{m:02d}"
    s = str(raw).strip()
    m = re.search(r'(\d{1,2})[:.](\d{2})', s)
    if m:
        h = int(m.group(1))
        mi = int(m.group(2))
        return f"{h:02d}:{mi:02d}"
    return None


def _extraer_numero_sala(raw: object) -> Optional[int]:
    """Extrae el numero entero de sala desde cualquier formato conocido.

    Casos reales observados en los Excel de Maritza:
        18.0        -> 18   (float de Excel)
        'SB-018'    -> 18   (prefijo SB con guion)
        'SB-O18'    -> 18   (typo: letra O en lugar de cero)
        'SB-14'     -> 14   (sin cero inicial)
        'SB -016'   -> 16   (espacio antes del guion)
        'CSC-020'   -> 20   (prefijo CSC edificio)
        'CSC - 020' -> 20   (prefijo CSC con espacios)
        '016'       -> 16   (solo numero con cero)
        'Sala 016'  -> 16   (formato canonico Hestia)
        16          -> 16   (int directo)
    """
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        num = int(raw)
        return num if num > 0 else None
    s = str(raw).strip()
    # Reemplazar letra O mayuscula por cero cuando aparece entre
    # prefijo no numerico y digitos (typo frecuente: SB-O18)
    s = re.sub(r'(?<=[A-Za-z\-\s])O(?=\d)', '0', s)
    grupos = re.findall(r'\d+', s)
    if not grupos:
        return None
    if len(grupos) == 1:
        return int(grupos[0])
    # Varios grupos: tomar el ultimo mayor que 0 (el numero de sala real)
    for g in reversed(grupos):
        n = int(g)
        if n > 0:
            return n
    return None


def _normalizar_sala_nombre(raw: object) -> str:
    """Convierte cualquier variante de nombre de sala al formato 'Sala NNN'."""
    num = _extraer_numero_sala(raw)
    if num is None:
        return str(raw).strip() if raw is not None else ""
    return f"Sala {num:03d}"


def _parsear_fecha(raw: object, semestre: str) -> Optional[dt.date]:
    """Resuelve una celda de fecha a datetime.date.

    Tipos soportados:
        datetime.datetime  -> .date()  (caso mas comun en openpyxl)
        datetime.date      -> directo
        str 'D-mes'        -> '4-sep', '10-sep', '24-sep' (texto en espanol)
        str 'DD-mes-AAAA'  -> '14-sep-2025'
        str 'AAAA-MM-DD'   -> ISO 8601

    Para formatos sin ano ('4-sep') se infiere el ano desde el semestre
    recibido ('2025-2' -> 2025).
    """
    if raw is None:
        return None

    if isinstance(raw, dt.datetime):
        return raw.date()
    if isinstance(raw, dt.date):
        return raw

    s = str(raw).strip()

    # Formato ISO: 2025-08-14
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})$', s)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))

    # Formato 'D-mes' o 'DD-mes' o 'D-mes-AAAA': '4-sep', '10-sep', '14-sep-2025'
    m = re.match(
        r'^(\d{1,2})[\-/\s](\w+?)(?:[\-/\s](\d{4}))?$', s, re.IGNORECASE
    )
    if m:
        dia = int(m.group(1))
        mes_s = m.group(2).lower()[:3]   # primeras 3 letras
        ano_s = m.group(3)

        mes = _MESES_ES.get(mes_s)
        if mes is None:
            return None

        if ano_s:
            ano = int(ano_s)
        else:
            # Inferir ano desde semestre (ej: '2025-2' -> 2025)
            try:
                ano = int(semestre.split('-')[0])
            except (ValueError, IndexError, AttributeError):
                ano = dt.date.today().year

        try:
            return dt.date(ano, mes, dia)
        except ValueError:
            return None

    return None


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@router.get("/hoy", response_model=list[ProgramacionTallerResponse])
def programacion_hoy(
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual),
):
    """Retorna todas las programaciones activas para la fecha de hoy."""
    hoy = date.today()
    rows = (
        db.query(ProgramacionTaller)
        .options(
            joinedload(ProgramacionTaller.taller).joinedload(
                Taller.asignatura
            ),
            joinedload(ProgramacionTaller.sala),
        )
        .filter(
            ProgramacionTaller.fecha == hoy,
            ProgramacionTaller.activo.is_(True),
        )
        .order_by(ProgramacionTaller.hora_inicio)
        .all()
    )
    return [_to_response(r) for r in rows]


@router.get("/", response_model=list[ProgramacionTallerResponse])
def listar_programacion(
    sala_id: Optional[int] = None,
    fecha: Optional[date] = None,
    semestre: Optional[str] = None,
    solo_activas: bool = True,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual),
):
    q = (
        db.query(ProgramacionTaller)
        .options(
            joinedload(ProgramacionTaller.taller).joinedload(
                Taller.asignatura
            ),
            joinedload(ProgramacionTaller.sala),
        )
    )
    if solo_activas:
        q = q.filter(ProgramacionTaller.activo.is_(True))
    if sala_id:
        q = q.filter(ProgramacionTaller.sala_id == sala_id)
    if fecha:
        q = q.filter(ProgramacionTaller.fecha == fecha)
    if semestre:
        q = q.filter(ProgramacionTaller.semestre == semestre)
    rows = (
        q.order_by(ProgramacionTaller.fecha, ProgramacionTaller.hora_inicio)
        .offset(skip).limit(limit).all()
    )
    return [_to_response(r) for r in rows]


@router.get("/{prog_id}", response_model=ProgramacionTallerResponse)
def obtener_programacion(
    prog_id: int,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual),
):
    return _to_response(_load(prog_id, db))


@router.post("/", response_model=ProgramacionTallerResponse, status_code=201)
def crear_programacion(
    datos: ProgramacionTallerCreate,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_operador),
):
    p = ProgramacionTaller(**datos.model_dump())
    db.add(p)
    db.commit()
    return _to_response(_load(p.id, db))


@router.put("/{prog_id}", response_model=ProgramacionTallerResponse)
def actualizar_programacion(
    prog_id: int,
    datos: ProgramacionTallerUpdate,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_operador),
):
    p = _load(prog_id, db)
    for campo, valor in datos.model_dump(exclude_unset=True).items():
        setattr(p, campo, valor)
    db.commit()
    return _to_response(_load(prog_id, db))


@router.delete("/{prog_id}", status_code=204)
def eliminar_programacion(
    prog_id: int,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_operador),
):
    p = _load(prog_id, db)
    p.activo = False
    db.commit()


@router.post(
    "/importar-xlsx",
    response_model=ImportarProgramacionResponse,
    status_code=200,
)
def importar_xlsx(
    semestre: str,
    archivo: UploadFile = File(...),
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_operador),
):
    """Importa programacion desde un xlsx de planificacion de Maritza.

    Estructura esperada de cada hoja:
        Columnas: GUIA DE TALLER | Nombre de Taller | Fecha | Semana
                  | Sala | Horario | Docente | Seccion

    Upsert por (taller_id, sala_id, fecha, seccion). Idempotente.
    """
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl no instalado en el servidor",
        )

    contenido = archivo.file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
    except Exception as exc:
        raise HTTPException(
            status_code=400, detail=f"Archivo xlsx invalido: {exc}"
        )

    salas_cache: dict[str, int] = {
        s.nombre: s.id for s in db.query(Sala).all()
    }
    talleres_cache: dict[str, int] = {
        t.nombre.strip().lower(): t.id
        for t in db.query(Taller).all()
    }

    importadas = actualizadas = omitidas = 0
    errores: list[dict] = []

    for hoja_nombre in wb.sheetnames:
        ws = wb[hoja_nombre]
        filas = list(ws.iter_rows(min_row=1, values_only=True))
        if not filas:
            continue

        header_idx = None
        for i, fila in enumerate(filas):
            fila_str = [str(v).strip().lower() if v else "" for v in fila]
            if "fecha" in fila_str:
                header_idx = i
                break
        if header_idx is None:
            omitidas += 1
            continue

        header = [
            str(v).strip().lower() if v else ""
            for v in filas[header_idx]
        ]

        def col(nombre_col: str) -> Optional[int]:
            try:
                return header.index(nombre_col)
            except ValueError:
                return None

        ci_taller = col("nombre de taller") or col("guia de taller")
        ci_fecha = col("fecha")
        ci_sala = col("sala")
        ci_horario = col("horario")
        ci_docente = col("docente")
        ci_seccion = col("seccion") or col("secci\u00f3n")

        if ci_fecha is None or ci_taller is None:
            errores.append({
                "hoja": hoja_nombre,
                "fila": header_idx + 1,
                "razon": "No se encontraron columnas Fecha o Nombre de Taller",
            })
            continue

        for fila_num, fila in enumerate(
            filas[header_idx + 1:], start=header_idx + 2
        ):
            if all(v is None for v in fila):
                continue

            raw_taller = fila[ci_taller] if ci_taller is not None else None
            raw_fecha = fila[ci_fecha]
            raw_sala = fila[ci_sala] if ci_sala is not None else None
            raw_horario = fila[ci_horario] if ci_horario is not None else None
            raw_docente = fila[ci_docente] if ci_docente is not None else None
            raw_seccion = fila[ci_seccion] if ci_seccion is not None else None

            if not raw_taller or not raw_fecha:
                omitidas += 1
                continue

            # Resolver fecha con ETL robusto
            fecha_val = _parsear_fecha(raw_fecha, semestre)
            if fecha_val is None:
                errores.append({
                    "hoja": hoja_nombre,
                    "fila": fila_num,
                    "razon": f"Fecha no reconocida: {raw_fecha!r}",
                })
                omitidas += 1
                continue

            # Resolver taller (buscar o crear)
            nombre_taller = str(raw_taller).strip()
            nombre_taller_key = nombre_taller.lower()
            if nombre_taller_key not in talleres_cache:
                nuevo_taller = Taller(nombre=nombre_taller)
                db.add(nuevo_taller)
                db.flush()
                talleres_cache[nombre_taller_key] = nuevo_taller.id
            taller_id = talleres_cache[nombre_taller_key]

            # Resolver sala con ETL robusto
            sala_id: Optional[int] = None
            if raw_sala is not None:
                nombre_sala_norm = _normalizar_sala_nombre(raw_sala)
                sala_id = salas_cache.get(nombre_sala_norm)
                if sala_id is None:
                    num_buscado = _extraer_numero_sala(raw_sala)
                    if num_buscado is not None:
                        for nombre_s, sid in salas_cache.items():
                            if _extraer_numero_sala(nombre_s) == num_buscado:
                                sala_id = sid
                                break

            if sala_id is None:
                errores.append({
                    "hoja": hoja_nombre,
                    "fila": fila_num,
                    "razon": f"Sala no encontrada: {raw_sala!r}",
                })
                omitidas += 1
                continue

            # Parsear horario
            hora_inicio: Optional[str] = None
            hora_fin: Optional[str] = None
            if raw_horario is not None:
                horario_str = str(raw_horario).strip()
                partes = re.split(
                    r'[\-aA]|\s+A\s+|\s+a\s+', horario_str
                )
                hora_inicio = _normalizar_hora(partes[0].strip())
                if len(partes) >= 2:
                    hora_fin = _normalizar_hora(partes[-1].strip())

            # Seccion
            seccion_val: Optional[str] = None
            if raw_seccion is not None:
                sv = str(raw_seccion).strip()
                if sv.endswith(".0"):
                    sv = sv[:-2]
                seccion_val = sv if sv else None

            docente_val = (
                str(raw_docente).strip() if raw_docente else None
            )

            # Upsert
            existente = (
                db.query(ProgramacionTaller)
                .filter(
                    and_(
                        ProgramacionTaller.taller_id == taller_id,
                        ProgramacionTaller.sala_id == sala_id,
                        ProgramacionTaller.fecha == fecha_val,
                        ProgramacionTaller.seccion == seccion_val,
                    )
                )
                .first()
            )

            if existente:
                existente.hora_inicio = hora_inicio
                existente.hora_fin = hora_fin
                existente.docente_nombre = docente_val
                existente.semestre = semestre
                existente.activo = True
                actualizadas += 1
            else:
                nueva = ProgramacionTaller(
                    taller_id=taller_id,
                    sala_id=sala_id,
                    fecha=fecha_val,
                    hora_inicio=hora_inicio,
                    hora_fin=hora_fin,
                    docente_nombre=docente_val,
                    seccion=seccion_val,
                    semestre=semestre,
                )
                db.add(nueva)
                importadas += 1

    db.commit()
    return ImportarProgramacionResponse(
        importadas=importadas,
        actualizadas=actualizadas,
        omitidas=omitidas,
        errores=errores,
    )
