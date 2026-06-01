from fastapi import APIRouter, Depends, HTTPException, Request, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import desc
from datetime import date, timedelta, datetime, timezone
from typing import Optional
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.models.movimiento import Movimiento, TipoMovimiento, SubtipoMovimiento
from app.models.insumo import Insumo, TipoInsumo
from app.models.usuario import Usuario
from app.models.sala import Sala
from app.models.retorno_implemento import RetornoImplemento
from app.schemas.movimiento import (
    MovimientoCreate,
    MovimientoResponse,
    MovimientoEnriquecido,
    EntregaDirectaCreate,
    EntregaDirectaResponse,
)
from app.schemas.comun import PaginatedResponse
from app.utils.deps import get_usuario_actual, require_operador
from app.utils.auditoria import registrar, get_ip

router = APIRouter(prefix="/movimientos", tags=["Movimientos"])

# Subtipos que reducen stock al registrarse
_SUBTIPOS_SALIDA = {
    SubtipoMovimiento.consumo_taller,
    SubtipoMovimiento.prestamo_implemento,
    SubtipoMovimiento.devolucion_proveedor_salida,
    SubtipoMovimiento.baja,
    SubtipoMovimiento.ajuste_salida,
}

# Subtipos que aumentan stock al registrarse
_SUBTIPOS_ENTRADA = {
    SubtipoMovimiento.compra,
    SubtipoMovimiento.devolucion_proveedor_entrada,
    SubtipoMovimiento.ajuste_entrada,
}

# Subtipos internos: no tocan stock (solo cambian estado del item)
_SUBTIPOS_INTERNOS = {
    SubtipoMovimiento.enviado_mantenimiento,
    SubtipoMovimiento.reingreso_disponible,
    SubtipoMovimiento.devolucion_interna,
}

# Tabla de validacion cruzada tipo <-> subtipo permitidos
_TIPO_SUBTIPO_VALIDO: dict[TipoMovimiento, set[SubtipoMovimiento]] = {
    TipoMovimiento.entrada: _SUBTIPOS_ENTRADA,
    TipoMovimiento.salida: _SUBTIPOS_SALIDA,
    TipoMovimiento.interno: _SUBTIPOS_INTERNOS,
}


def _validar_tipo_subtipo(
    tipo: TipoMovimiento, subtipo: SubtipoMovimiento
) -> None:
    """Lanza 422 si el subtipo no corresponde al tipo base."""
    permitidos = _TIPO_SUBTIPO_VALIDO.get(tipo, set())
    if subtipo not in permitidos:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                f"El subtipo '{subtipo.value}' no es valido "
                f"para el tipo '{tipo.value}'."
            ),
        )


def _enriquecer(m: Movimiento) -> MovimientoEnriquecido:
    return MovimientoEnriquecido(
        id=m.id,
        tipo=m.tipo,
        subtipo=m.subtipo,
        cantidad=m.cantidad,
        motivo=m.motivo,
        fecha=m.fecha,
        insumo=m.insumo.nombre if m.insumo else "Desconocido",
        sala=m.sala.nombre if m.sala else None,
        usuario=m.usuario.nombre if m.usuario else "Desconocido",
        paquete_id=m.paquete_id,
    )


def _build_query(
    db: Session,
    insumo: Optional[str] = None,
    tipo: Optional[str] = None,
    subtipo: Optional[str] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
):
    """Query base con filtros opcionales para listado y exportacion.

    fecha_hasta es inclusivo: se suma un dia para incluir todo el dia UTC.
    """
    q = (
        db.query(Movimiento)
        .options(
            joinedload(Movimiento.insumo),
            joinedload(Movimiento.usuario),
            joinedload(Movimiento.sala),
        )
        .order_by(desc(Movimiento.fecha))
    )
    if insumo:
        ids = (
            db.query(Insumo.id)
            .filter(Insumo.nombre.ilike(f"%{insumo}%"))
            .subquery()
        )
        q = q.filter(Movimiento.insumo_id.in_(ids))
    if tipo and tipo in ("entrada", "salida", "interno"):
        q = q.filter(Movimiento.tipo == tipo)
    if subtipo:
        try:
            subtipo_enum = SubtipoMovimiento(subtipo)
            q = q.filter(Movimiento.subtipo == subtipo_enum)
        except ValueError:
            pass
    if fecha_desde:
        q = q.filter(Movimiento.fecha >= fecha_desde)
    if fecha_hasta:
        q = q.filter(Movimiento.fecha < fecha_hasta + timedelta(days=1))
    return q


# ---------------------------------------------------------------------------
# IMPORTANTE: /exportar y /entrega-directa deben ir ANTES de las rutas
# dinamicas /{id}, de lo contrario FastAPI los interpreta como enteros.
# ---------------------------------------------------------------------------

@router.get("/exportar")
def exportar_movimientos(
    formato: str = "csv",
    insumo: Optional[str] = None,
    tipo: Optional[str] = None,
    subtipo: Optional[str] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual),
):
    """Exporta movimientos como CSV o Excel con los mismos filtros del listado.

    formato=csv  -> archivo .csv con BOM UTF-8 (abre bien en Excel).
    formato=xlsx -> archivo .xlsx con cabecera coloreada y columnas autoajustadas.
    """
    movimientos = _build_query(
        db, insumo, tipo, subtipo, fecha_desde, fecha_hasta
    ).all()
    filas = [
        [
            m.tipo.value if hasattr(m.tipo, "value") else m.tipo,
            m.subtipo.value if m.subtipo and hasattr(m.subtipo, "value")
            else (m.subtipo or ""),
            m.insumo.nombre if m.insumo else "Desconocido",
            m.sala.nombre if m.sala else "",
            m.cantidad,
            m.motivo or "",
            m.fecha.strftime("%d/%m/%Y %H:%M") if m.fecha else "",
            m.usuario.nombre if m.usuario else "Desconocido",
        ]
        for m in movimientos
    ]
    cabecera = [
        "Tipo", "Subtipo", "Insumo", "Sala",
        "Cantidad", "Motivo", "Fecha", "Usuario",
    ]

    if formato == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Movimientos"

        header_fill = PatternFill("solid", fgColor="0F766E")  # teal-700
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx, titulo in enumerate(cabecera, 1):
            cell = ws.cell(row=1, column=col_idx, value=titulo)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for fila in filas:
            ws.append(fila)

        for col_idx, _ in enumerate(cabecera, 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                (
                    len(str(ws.cell(row=r, column=col_idx).value or ""))
                    for r in range(1, ws.max_row + 1)
                ),
                default=10,
            )
            ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type=(
                "application/vnd.openxmlformats-officedocument"
                ".spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                    "attachment; filename=movimientos_hestia.xlsx"
            },
        )

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(cabecera)
    writer.writerows(filas)
    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={
            "Content-Disposition":
                "attachment; filename=movimientos_hestia.csv"
        },
    )


@router.post("/entrega-directa", response_model=EntregaDirectaResponse)
def entrega_directa(
    request: Request,
    datos: EntregaDirectaCreate,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_operador),
):
    """Operador registra una entrega inmediata a un docente presente.

    Usa subtipo=prestamo_implemento para implementos y
    subtipo=consumo_taller para insumos desechables.
    """
    sala = db.query(Sala).filter(Sala.id == datos.sala_id).first()
    if not sala:
        raise HTTPException(status_code=404, detail="Sala no encontrada")

    docente = db.query(Usuario).filter(
        Usuario.id == datos.docente_id,
        Usuario.activo.is_(True),
    ).first()
    if not docente:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="El usuario indicado no existe o esta inactivo",
        )

    cantidades: dict[int, int] = {}
    for item in datos.items:
        cantidades[item.insumo_id] = (
            cantidades.get(item.insumo_id, 0) + item.cantidad
        )

    insumos_bloqueados: dict[int, Insumo] = {}
    for insumo_id, cantidad in cantidades.items():
        insumo = (
            db.query(Insumo)
            .filter(Insumo.id == insumo_id)
            .with_for_update()
            .first()
        )
        nombre = insumo.nombre if insumo else str(insumo_id)
        if not insumo or not insumo.activo:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f"'{nombre}' no esta disponible o fue desactivado",
            )
        if insumo.stock_actual < cantidad:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=(
                    f"Stock insuficiente para '{insumo.nombre}'. "
                    f"Disponible: {insumo.stock_actual}, "
                    f"solicitado: {cantidad}."
                ),
            )
        insumos_bloqueados[insumo_id] = insumo

    motivo = f"Entrega directa \u2014 {docente.nombre} \u2014 {sala.nombre}"
    ahora = datetime.now(timezone.utc)
    retornos_generados = 0

    for insumo_id, cantidad in cantidades.items():
        insumo = insumos_bloqueados[insumo_id]
        insumo.stock_actual -= cantidad

        subtipo = (
            SubtipoMovimiento.prestamo_implemento
            if insumo.tipo == TipoInsumo.implemento
            else SubtipoMovimiento.consumo_taller
        )
        db.add(Movimiento(
            tipo=TipoMovimiento.salida,
            subtipo=subtipo,
            cantidad=cantidad,
            insumo_id=insumo_id,
            usuario_id=usuario.id,
            sala_id=datos.sala_id,
            motivo=motivo,
        ))
        if insumo.tipo == TipoInsumo.implemento:
            db.add(RetornoImplemento(
                insumo_id=insumo_id,
                solicitud_id=None,
                docente_id=datos.docente_id,
                sala_id=datos.sala_id,
                cantidad=cantidad,
                fecha_retiro=ahora,
            ))
            retornos_generados += 1

    db.commit()
    registrar(
        db,
        "ENTREGA_DIRECTA",
        usuario=usuario,
        entidad="movimiento",
        entidad_id=None,
        detalle=(
            f"{len(cantidades)} item(s) a {docente.nombre}"
            f" en {sala.nombre}"
        ),
        ip=get_ip(request),
    )
    return EntregaDirectaResponse(
        mensaje="Entrega registrada correctamente",
        items_procesados=len(cantidades),
        retornos_pendientes=retornos_generados,
    )


@router.get("/", response_model=PaginatedResponse[MovimientoEnriquecido])
def listar_movimientos(
    skip: int = 0,
    limit: int = 20,
    insumo: Optional[str] = None,
    tipo: Optional[str] = None,
    subtipo: Optional[str] = None,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual),
):
    """Lista movimientos con filtros opcionales: insumo (texto), tipo,
    subtipo, fecha_desde y fecha_hasta. Todos los filtros son acumulables.
    """
    q = _build_query(db, insumo, tipo, subtipo, fecha_desde, fecha_hasta)
    total = q.count()
    movimientos = q.offset(skip).limit(limit).all()
    return {
        "total": total,
        "skip": skip,
        "limit": limit,
        "data": [_enriquecer(m) for m in movimientos],
    }


@router.post("/", response_model=MovimientoResponse)
def registrar_movimiento(
    request: Request,
    mov: MovimientoCreate,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(require_operador),
):
    """Registra un movimiento validando la coherencia tipo<->subtipo.

    Regla de stock:
    - ENTRADA: aumenta stock_actual
    - SALIDA:  reduce stock_actual (valida suficiencia)
    - INTERNO: no toca stock (solo trazabilidad de estado)
    """
    _validar_tipo_subtipo(mov.tipo, mov.subtipo)

    insumo = (
        db.query(Insumo)
        .filter(Insumo.id == mov.insumo_id)
        .with_for_update()
        .first()
    )
    if not insumo:
        raise HTTPException(status_code=404, detail="Insumo no encontrado")

    if not insumo.activo:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=(
                f"'{insumo.nombre}' esta inactivo. Reactivalo antes de "
                "registrar movimientos."
            ),
        )

    if mov.tipo == TipoMovimiento.salida:
        if insumo.stock_actual < mov.cantidad:
            raise HTTPException(
                status_code=400,
                detail=f"Stock insuficiente. Disponible: {insumo.stock_actual}",
            )
        insumo.stock_actual -= mov.cantidad
    elif mov.tipo == TipoMovimiento.entrada:
        insumo.stock_actual += mov.cantidad
    # interno: sin cambio de stock

    nuevo_mov = Movimiento(
        **mov.model_dump(),
        usuario_id=usuario.id,
    )
    db.add(nuevo_mov)
    db.commit()
    db.refresh(nuevo_mov)
    registrar(
        db,
        "REGISTRAR_MOVIMIENTO",
        usuario=usuario,
        entidad="movimiento",
        entidad_id=nuevo_mov.id,
        detalle=(
            f"{mov.tipo.value}/{mov.subtipo.value} "
            f"{mov.cantidad}x {insumo.nombre}"
        ),
        ip=get_ip(request),
    )
    return nuevo_mov


@router.get("/insumo/{insumo_id}", response_model=PaginatedResponse[MovimientoEnriquecido])
def historial_por_insumo(
    insumo_id: int,
    skip: int = 0,
    limit: int = 20,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual),
):
    if not db.query(Insumo).filter(Insumo.id == insumo_id).first():
        raise HTTPException(status_code=404, detail="Insumo no encontrado")
    total = db.query(Movimiento).filter(
        Movimiento.insumo_id == insumo_id
    ).count()
    movimientos = (
        _build_query(db)
        .filter(Movimiento.insumo_id == insumo_id)
        .offset(skip).limit(limit).all()
    )
    return {
        "total": total, "skip": skip, "limit": limit,
        "data": [_enriquecer(m) for m in movimientos],
    }


@router.get("/sala/{sala_id}", response_model=PaginatedResponse[MovimientoEnriquecido])
def historial_por_sala(
    sala_id: int,
    skip: int = 0,
    limit: int = 20,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(get_usuario_actual),
):
    if not db.query(Sala).filter(Sala.id == sala_id).first():
        raise HTTPException(status_code=404, detail="Sala no encontrada")
    total = (
        db.query(Movimiento)
        .filter(Movimiento.sala_id == sala_id)
        .count()
    )
    movimientos = (
        _build_query(db)
        .filter(Movimiento.sala_id == sala_id)
        .offset(skip).limit(limit).all()
    )
    return {
        "total": total, "skip": skip, "limit": limit,
        "data": [_enriquecer(m) for m in movimientos],
    }
