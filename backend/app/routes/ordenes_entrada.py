from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Request, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session, joinedload

from app.database import get_db
from app.models.orden_entrada import (
    OrdenEntrada, OrdenEntradaItem,
    EstadoOrden, EstadoItem, TipoItemOrden, ACTIVIDADES_DUOC,
)
from app.models.insumo import Insumo, TipoInsumo
from app.models.activo_fijo import ActivoFijo
from app.models.movimiento import Movimiento, TipoMovimiento, SubtipoMovimiento
from app.models.usuario import Usuario, RolUsuario
from app.schemas.orden_entrada import (
    OrdenEntradaCreate, OrdenEntradaUpdate,
    OrdenEntradaItemCreate, OrdenEntradaItemUpdate,
    OrdenEntradaItemResponse, OrdenEntradaResponse,
)
from app.utils.deps import get_usuario_actual, require_operador
from app.utils.auditoria import registrar, get_ip

router = APIRouter(prefix="/ordenes-entrada", tags=["Ordenes de Entrada"])


def _require_coord_o_admin(usuario: Usuario = Depends(get_usuario_actual)):
    if usuario.rol not in (RolUsuario.admin, RolUsuario.operador_coordinador):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Solo el coordinador o administrador puede gestionar ordenes.",
        )
    return usuario


def _to_response(o: OrdenEntrada) -> OrdenEntradaResponse:
    items_resp = []
    for i in o.items:
        items_resp.append(
            OrdenEntradaItemResponse(
                id=i.id,
                orden_id=i.orden_id,
                tipo_item=(
                    i.tipo_item.value if hasattr(i.tipo_item, 'value')
                    else i.tipo_item
                ),
                insumo_id=i.insumo_id,
                insumo_nombre=i.insumo.nombre if i.insumo else None,
                activo_fijo_id=i.activo_fijo_id,
                activo_fijo_nombre=(
                    i.activo_fijo.nombre if i.activo_fijo else None
                ),
                nombre_nuevo=i.nombre_nuevo,
                tipo_insumo_nuevo=i.tipo_insumo_nuevo,
                tipo_activo_nuevo=i.tipo_activo_nuevo,
                cantidad_pedida=i.cantidad_pedida,
                cantidad_recibida=i.cantidad_recibida,
                costo_unitario=(
                    float(i.costo_unitario) if i.costo_unitario else None
                ),
                estado=(
                    i.estado.value if hasattr(i.estado, 'value')
                    else i.estado
                ),
                notas_item=i.notas_item,
            )
        )
    total_pedido = sum(i.cantidad_pedida for i in o.items)
    total_recibido = sum(
        i.cantidad_recibida
        for i in o.items
        if i.cantidad_recibida is not None
    )
    act_nombre = (
        ACTIVIDADES_DUOC.get(o.actividad_duoc) if o.actividad_duoc else None
    )
    return OrdenEntradaResponse(
        id=o.id,
        proveedor_id=o.proveedor_id,
        proveedor_nombre=o.proveedor.nombre if o.proveedor else None,
        actividad_duoc=o.actividad_duoc,
        actividad_nombre=act_nombre,
        tipo=o.tipo.value if hasattr(o.tipo, 'value') else o.tipo,
        estado=o.estado.value if hasattr(o.estado, 'value') else o.estado,
        notas=o.notas,
        creado_por_id=o.creado_por_id,
        creado_por_nombre=o.creado_por.nombre if o.creado_por else None,
        cerrado_por_id=o.cerrado_por_id,
        cerrado_por_nombre=o.cerrado_por.nombre if o.cerrado_por else None,
        created_at=o.created_at,
        fecha_cierre=o.fecha_cierre,
        items=items_resp,
        total_pedido=total_pedido,
        total_recibido=total_recibido,
    )


def _cargar(orden_id: int, db: Session) -> OrdenEntrada:
    o = (
        db.query(OrdenEntrada)
        .options(
            joinedload(OrdenEntrada.proveedor),
            joinedload(OrdenEntrada.creado_por),
            joinedload(OrdenEntrada.cerrado_por),
            joinedload(OrdenEntrada.items)
            .joinedload(OrdenEntradaItem.insumo),
            joinedload(OrdenEntrada.items)
            .joinedload(OrdenEntradaItem.activo_fijo),
        )
        .filter(OrdenEntrada.id == orden_id)
        .first()
    )
    if not o:
        raise HTTPException(status_code=404, detail="Orden no encontrada.")
    return o


def _make_item(orden_id: int, datos: OrdenEntradaItemCreate) -> OrdenEntradaItem:
    return OrdenEntradaItem(
        orden_id=orden_id,
        tipo_item=datos.tipo_item,
        insumo_id=datos.insumo_id,
        activo_fijo_id=datos.activo_fijo_id,
        nombre_nuevo=datos.nombre_nuevo,
        tipo_insumo_nuevo=datos.tipo_insumo_nuevo,
        tipo_activo_nuevo=datos.tipo_activo_nuevo,
        cantidad_pedida=datos.cantidad_pedida,
        costo_unitario=datos.costo_unitario,
        notas_item=datos.notas_item,
        estado=EstadoItem.pendiente,
    )


# ---------------------------------------------------------------------------
# Endpoints
# IMPORTANTE: rutas estaticas antes de rutas dinamicas /{orden_id}
# ---------------------------------------------------------------------------

@router.get("/actividades", response_model=list[dict])
def listar_actividades(_=Depends(get_usuario_actual)):
    return [{"codigo": k, "nombre": v} for k, v in ACTIVIDADES_DUOC.items()]


@router.get("/", response_model=list[OrdenEntradaResponse])
def listar(
    estado: Optional[str] = None,
    proveedor_id: Optional[int] = None,
    skip: int = 0,
    limit: int = 50,
    db: Session = Depends(get_db),
    _=Depends(get_usuario_actual),
):
    q = (
        db.query(OrdenEntrada)
        .options(
            joinedload(OrdenEntrada.proveedor),
            joinedload(OrdenEntrada.creado_por),
            joinedload(OrdenEntrada.cerrado_por),
            joinedload(OrdenEntrada.items)
            .joinedload(OrdenEntradaItem.insumo),
            joinedload(OrdenEntrada.items)
            .joinedload(OrdenEntradaItem.activo_fijo),
        )
    )
    if estado:
        q = q.filter(OrdenEntrada.estado == estado)
    if proveedor_id:
        q = q.filter(OrdenEntrada.proveedor_id == proveedor_id)
    ordenes = (
        q.order_by(OrdenEntrada.created_at.desc())
        .offset(skip).limit(limit).all()
    )
    return [_to_response(o) for o in ordenes]


@router.post("/", response_model=OrdenEntradaResponse, status_code=201)
def crear(
    request: Request,
    datos: OrdenEntradaCreate,
    db: Session = Depends(get_db),
    usuario=Depends(_require_coord_o_admin),
):
    o = OrdenEntrada(
        proveedor_id=datos.proveedor_id,
        actividad_duoc=datos.actividad_duoc,
        tipo=datos.tipo,
        estado=EstadoOrden.borrador,
        notas=datos.notas,
        creado_por_id=usuario.id,
    )
    db.add(o)
    db.flush()
    for it in datos.items:
        db.add(_make_item(o.id, it))
    db.commit()
    registrar(
        db, "CREAR_ORDEN_ENTRADA", usuario=usuario,
        entidad="orden_entrada", entidad_id=o.id,
        detalle=f"Tipo {datos.tipo} - {len(datos.items)} items",
        ip=get_ip(request),
    )
    return _to_response(_cargar(o.id, db))


@router.get("/{orden_id}", response_model=OrdenEntradaResponse)
def obtener(
    orden_id: int,
    db: Session = Depends(get_db),
    _=Depends(get_usuario_actual),
):
    return _to_response(_cargar(orden_id, db))


@router.put("/{orden_id}", response_model=OrdenEntradaResponse)
def actualizar(
    orden_id: int,
    request: Request,
    datos: OrdenEntradaUpdate,
    db: Session = Depends(get_db),
    usuario=Depends(_require_coord_o_admin),
):
    o = _cargar(orden_id, db)
    if o.estado != EstadoOrden.borrador:
        raise HTTPException(
            status_code=400,
            detail="Solo se puede editar una orden en estado borrador.",
        )
    for campo, valor in datos.model_dump(exclude_unset=True).items():
        setattr(o, campo, valor)
    db.commit()
    return _to_response(_cargar(orden_id, db))


@router.post("/{orden_id}/items", response_model=OrdenEntradaResponse)
def agregar_item(
    orden_id: int,
    datos: OrdenEntradaItemCreate,
    db: Session = Depends(get_db),
    usuario=Depends(_require_coord_o_admin),
):
    o = _cargar(orden_id, db)
    if o.estado != EstadoOrden.borrador:
        raise HTTPException(
            status_code=400, detail="La orden no esta en borrador."
        )
    db.add(_make_item(orden_id, datos))
    db.commit()
    return _to_response(_cargar(orden_id, db))


class BulkItemsPayload(BaseModel):
    items: list[OrdenEntradaItemCreate]


@router.post(
    "/{orden_id}/items/bulk",
    response_model=OrdenEntradaResponse,
)
def agregar_items_bulk(
    orden_id: int,
    payload: BulkItemsPayload,
    db: Session = Depends(get_db),
    usuario=Depends(_require_coord_o_admin),
):
    """Agrega multiples items en una sola transaccion atomica.

    Si cualquier item es invalido, ningun item se guarda.
    Requiere que la orden este en estado borrador.
    """
    o = _cargar(orden_id, db)
    if o.estado != EstadoOrden.borrador:
        raise HTTPException(
            status_code=400, detail="La orden no esta en borrador."
        )
    if not payload.items:
        raise HTTPException(
            status_code=400, detail="Se requiere al menos un item."
        )
    for datos in payload.items:
        db.add(_make_item(orden_id, datos))
    db.commit()
    return _to_response(_cargar(orden_id, db))


@router.delete(
    "/{orden_id}/items/{item_id}",
    response_model=OrdenEntradaResponse,
)
def eliminar_item(
    orden_id: int,
    item_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(_require_coord_o_admin),
):
    o = _cargar(orden_id, db)
    if o.estado != EstadoOrden.borrador:
        raise HTTPException(
            status_code=400, detail="La orden no esta en borrador."
        )
    item = db.query(OrdenEntradaItem).filter(
        OrdenEntradaItem.id == item_id,
        OrdenEntradaItem.orden_id == orden_id,
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item no encontrado.")
    db.delete(item)
    db.commit()
    return _to_response(_cargar(orden_id, db))


@router.post("/{orden_id}/confirmar", response_model=OrdenEntradaResponse)
def confirmar(
    orden_id: int,
    request: Request,
    db: Session = Depends(get_db),
    usuario=Depends(_require_coord_o_admin),
):
    o = _cargar(orden_id, db)
    if o.estado != EstadoOrden.borrador:
        raise HTTPException(
            status_code=400, detail="La orden ya fue confirmada."
        )
    if not o.items:
        raise HTTPException(
            status_code=400, detail="La orden no tiene items."
        )
    o.estado = EstadoOrden.confirmada
    db.commit()
    registrar(
        db, "CONFIRMAR_ORDEN_ENTRADA", usuario=usuario,
        entidad="orden_entrada", entidad_id=orden_id,
        detalle=(
            f"Proveedor: "
            f"{o.proveedor.nombre if o.proveedor else 'sin proveedor'}"
        ),
        ip=get_ip(request),
    )
    return _to_response(_cargar(orden_id, db))


@router.patch(
    "/{orden_id}/items/{item_id}/recepcion",
    response_model=OrdenEntradaResponse,
)
def registrar_recepcion_item(
    orden_id: int,
    item_id: int,
    datos: OrdenEntradaItemUpdate,
    db: Session = Depends(get_db),
    usuario=Depends(require_operador),
):
    """Operador registra cantidad fisicamente recibida por item."""
    o = _cargar(orden_id, db)
    if o.estado not in (EstadoOrden.confirmada, EstadoOrden.en_recepcion):
        raise HTTPException(
            status_code=400,
            detail="La orden debe estar confirmada para registrar recepcion.",
        )
    item = db.query(OrdenEntradaItem).filter(
        OrdenEntradaItem.id == item_id,
        OrdenEntradaItem.orden_id == orden_id,
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Item no encontrado.")
    if datos.cantidad_recibida is not None:
        item.cantidad_recibida = datos.cantidad_recibida
        if datos.cantidad_recibida == 0:
            item.estado = EstadoItem.cancelado
        elif datos.cantidad_recibida < item.cantidad_pedida:
            item.estado = EstadoItem.recibido_parcial
        else:
            item.estado = EstadoItem.recibido
    if datos.costo_unitario is not None:
        item.costo_unitario = datos.costo_unitario
    if datos.notas_item is not None:
        item.notas_item = datos.notas_item
    if o.estado == EstadoOrden.confirmada:
        o.estado = EstadoOrden.en_recepcion
    db.commit()
    return _to_response(_cargar(orden_id, db))


@router.post("/{orden_id}/cerrar", response_model=OrdenEntradaResponse)
def cerrar(
    orden_id: int,
    request: Request,
    db: Session = Depends(get_db),
    usuario=Depends(_require_coord_o_admin),
):
    """Cierra la orden: impacta stock, crea registros nuevos, genera movimientos."""
    o = _cargar(orden_id, db)
    if o.estado not in (EstadoOrden.confirmada, EstadoOrden.en_recepcion):
        raise HTTPException(
            status_code=400,
            detail="Solo se puede cerrar una orden confirmada o en recepcion.",
        )

    items_procesados = 0

    for item in o.items:
        cant = item.cantidad_recibida
        if not cant or cant <= 0:
            item.estado = EstadoItem.cancelado
            continue

        if item.tipo_item == TipoItemOrden.insumo:
            insumo = item.insumo
            if insumo is None and item.nombre_nuevo:
                tipo_nuevo = TipoInsumo(item.tipo_insumo_nuevo or "insumo")
                insumo = Insumo(
                    nombre=item.nombre_nuevo,
                    stock_actual=0,
                    stock_minimo=0,
                    tipo=tipo_nuevo,
                    activo=True,
                )
                db.add(insumo)
                db.flush()
                if not insumo.sku:
                    insumo.sku = f"HST-{insumo.id:05d}"
                item.insumo_id = insumo.id
                db.flush()
            if insumo:
                insumo.stock_actual += cant
                motivo = f"Orden entrada #{o.id}"
                if o.proveedor:
                    motivo += f" - {o.proveedor.nombre}"
                db.add(Movimiento(
                    tipo=TipoMovimiento.entrada,
                    subtipo=SubtipoMovimiento.compra,
                    cantidad=cant,
                    insumo_id=insumo.id,
                    usuario_id=usuario.id,
                    motivo=motivo,
                ))
                item.estado = (
                    EstadoItem.recibido
                    if cant >= item.cantidad_pedida
                    else EstadoItem.recibido_parcial
                )
                items_procesados += 1

        elif item.tipo_item == TipoItemOrden.activo_fijo:
            from app.models.activo_fijo import TipoActivo
            activo = item.activo_fijo
            if activo is None and item.nombre_nuevo:
                tipo_af = TipoActivo(item.tipo_activo_nuevo or "mueble")
                activo = ActivoFijo(
                    nombre=item.nombre_nuevo,
                    tipo=tipo_af,
                    activo=True,
                    proveedor_id=o.proveedor_id,
                )
                db.add(activo)
                db.flush()
                prefijo = "MUE" if tipo_af.value == "mueble" else "PHN"
                activo.codigo_interno = f"{prefijo}-{activo.id:05d}"
                item.activo_fijo_id = activo.id
                db.flush()
            item.estado = EstadoItem.recibido
            items_procesados += 1

    o.estado = EstadoOrden.cerrada
    o.cerrado_por_id = usuario.id
    o.fecha_cierre = datetime.now(timezone.utc)
    db.commit()

    detalle = f"{items_procesados} items procesados"
    if o.proveedor:
        detalle += f" - {o.proveedor.nombre}"
    registrar(
        db, "CERRAR_ORDEN_ENTRADA", usuario=usuario,
        entidad="orden_entrada", entidad_id=orden_id,
        detalle=detalle, ip=get_ip(request),
    )
    return _to_response(_cargar(orden_id, db))


@router.post("/{orden_id}/cancelar", response_model=OrdenEntradaResponse)
def cancelar(
    orden_id: int,
    request: Request,
    db: Session = Depends(get_db),
    usuario=Depends(_require_coord_o_admin),
):
    o = _cargar(orden_id, db)
    if o.estado == EstadoOrden.cerrada:
        raise HTTPException(
            status_code=400,
            detail="No se puede cancelar una orden cerrada.",
        )
    o.estado = EstadoOrden.cancelada
    db.commit()
    registrar(
        db, "CANCELAR_ORDEN_ENTRADA", usuario=usuario,
        entidad="orden_entrada", entidad_id=orden_id,
        detalle="Orden cancelada", ip=get_ip(request),
    )
    return _to_response(_cargar(orden_id, db))


@router.get("/{orden_id}/exportar-pdf")
def exportar_pdf(
    orden_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(get_usuario_actual),
):
    """Genera PDF. Solo disponible para ordenes en borrador, confirmada o cerrada."""
    try:
        from fpdf import FPDF
    except ImportError:
        raise HTTPException(status_code=500, detail="fpdf2 no instalado.")
    import io

    o = _cargar(orden_id, db)
    estados_permitidos = (
        EstadoOrden.borrador, EstadoOrden.confirmada, EstadoOrden.cerrada
    )
    if o.estado not in estados_permitidos:
        raise HTTPException(
            status_code=400,
            detail="El PDF solo esta disponible para ordenes en borrador, "
                   "confirmada o cerrada.",
        )

    pdf = FPDF()
    pdf.set_margins(15, 15, 15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Hestia - Orden de Entrada", ln=True)
    pdf.set_font("Helvetica", "", 10)
    estado_str = o.estado.value if hasattr(o.estado, 'value') else o.estado
    tipo_str = o.tipo.value if hasattr(o.tipo, 'value') else o.tipo
    pdf.cell(0, 6, f"N# {o.id}  |  Estado: {estado_str}", ln=True)
    pdf.cell(0, 6, f"Tipo: {tipo_str}", ln=True)
    if o.proveedor:
        pdf.cell(0, 6, f"Proveedor: {o.proveedor.nombre}", ln=True)
    if o.actividad_duoc:
        nombre_act = ACTIVIDADES_DUOC.get(o.actividad_duoc, o.actividad_duoc)
        pdf.cell(
            0, 6,
            f"Actividad DuocUC: ({o.actividad_duoc}) {nombre_act}",
            ln=True,
        )
    pdf.cell(
        0, 6, f"Creado: {o.created_at.strftime('%d/%m/%Y %H:%M')}", ln=True
    )
    if o.fecha_cierre:
        pdf.cell(
            0, 6,
            f"Cerrado: {o.fecha_cierre.strftime('%d/%m/%Y %H:%M')}",
            ln=True,
        )
    if o.notas:
        pdf.cell(0, 6, f"Notas: {o.notas}", ln=True)
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(15, 118, 110)
    pdf.set_text_color(255, 255, 255)
    cols = [
        ("Item", 65), ("Tipo", 22), ("Pedido", 20),
        ("Recibido", 22), ("Costo unit.", 28), ("Subtotal", 28),
    ]
    for txt, w in cols:
        pdf.cell(w, 7, txt, border=1, fill=True)
    pdf.ln()
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(0, 0, 0)

    total_costo = 0.0
    for it in o.items:
        nombre = (
            it.insumo.nombre if it.insumo
            else it.activo_fijo.nombre if it.activo_fijo
            else (it.nombre_nuevo or "Nuevo")
        )
        it_tipo = (
            it.tipo_item.value if hasattr(it.tipo_item, 'value')
            else it.tipo_item
        )
        recibido_str = (
            str(it.cantidad_recibida)
            if it.cantidad_recibida is not None else "-"
        )
        costo_str = (
            f"${float(it.costo_unitario):,.0f}" if it.costo_unitario else "-"
        )
        subtotal = (
            float(it.costo_unitario)
            * (it.cantidad_recibida or it.cantidad_pedida)
            if it.costo_unitario else 0
        )
        total_costo += subtotal
        subtotal_str = f"${subtotal:,.0f}" if it.costo_unitario else "-"
        for txt, w in [
            (nombre[:38], 65), (it_tipo, 22), (str(it.cantidad_pedida), 20),
            (recibido_str, 22), (costo_str, 28), (subtotal_str, 28),
        ]:
            pdf.cell(w, 6, txt, border=1)
        pdf.ln()

    if total_costo > 0:
        pdf.set_font("Helvetica", "B", 9)
        pdf.cell(137, 7, "TOTAL", border=1)
        pdf.cell(28, 7, f"${total_costo:,.0f}", border=1)
        pdf.ln()

    pdf.ln(8)
    pdf.set_font("Helvetica", "I", 7)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(
        0, 5,
        "Documento generado por Hestia - Escuela de Salud DuocUC San Bernardo",
        align="C",
    )

    buf = io.BytesIO(pdf.output())
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f"attachment; filename=orden_entrada_{o.id}.pdf"
        },
    )


@router.get("/{orden_id}/exportar-excel")
def exportar_excel(
    orden_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(get_usuario_actual),
):
    """Genera Excel. Solo para ordenes en borrador, confirmada o cerrada."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl no instalado.")
    import io

    o = _cargar(orden_id, db)
    estados_permitidos = (
        EstadoOrden.borrador, EstadoOrden.confirmada, EstadoOrden.cerrada
    )
    if o.estado not in estados_permitidos:
        raise HTTPException(
            status_code=400,
            detail="El Excel solo esta disponible para ordenes en borrador, "
                   "confirmada o cerrada.",
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Orden {o.id}"

    estado_str = o.estado.value if hasattr(o.estado, 'value') else o.estado
    ws.append(["Hestia - Orden de Entrada"])
    ws.append([f"N# {o.id}", f"Estado: {estado_str}"])
    ws.append(["Proveedor:", o.proveedor.nombre if o.proveedor else "-"])
    if o.actividad_duoc:
        act_nombre = ACTIVIDADES_DUOC.get(o.actividad_duoc, '')
        ws.append([
            "Actividad DuocUC:",
            f"({o.actividad_duoc}) {act_nombre}",
        ])
    ws.append(["Fecha:", o.created_at.strftime("%d/%m/%Y %H:%M")])
    ws.append([])

    hdrs = [
        "Item", "Tipo", "Es nuevo", "Pedido", "Recibido",
        "Costo unitario", "Subtotal", "Estado", "Notas",
    ]
    ws.append(hdrs)
    hdr_fill = PatternFill("solid", fgColor="0F766E")
    bold_white = Font(color="FFFFFF", bold=True)
    for col, _ in enumerate(hdrs, 1):
        cell = ws.cell(row=ws.max_row, column=col)
        cell.fill = hdr_fill
        cell.font = bold_white
        cell.alignment = Alignment(horizontal="center")

    total_costo = 0.0
    for it in o.items:
        nombre = (
            it.insumo.nombre if it.insumo
            else it.activo_fijo.nombre if it.activo_fijo
            else (it.nombre_nuevo or "Nuevo")
        )
        es_nuevo = (
            "Si" if (not it.insumo_id and not it.activo_fijo_id) else "No"
        )
        it_tipo = (
            it.tipo_item.value if hasattr(it.tipo_item, 'value')
            else str(it.tipo_item)
        )
        costo = float(it.costo_unitario) if it.costo_unitario else None
        recibido = it.cantidad_recibida
        subtotal = costo * (recibido or it.cantidad_pedida) if costo else None
        if subtotal:
            total_costo += subtotal
        ws.append([
            nombre, it_tipo, es_nuevo,
            it.cantidad_pedida, recibido, costo, subtotal,
            it.estado.value if hasattr(it.estado, 'value') else str(it.estado),
            it.notas_item or "",
        ])

    # Fila de total
    if total_costo > 0:
        ws.append(["", "", "", "", "", "TOTAL", total_costo, "", ""])
        total_row = ws.max_row
        for col in (6, 7):
            cell = ws.cell(row=total_row, column=col)
            cell.font = Font(bold=True)

    for col_idx in range(1, len(hdrs) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            (
                len(str(ws.cell(row=r, column=col_idx).value or ""))
                for r in range(1, ws.max_row + 1)
            ),
            default=10,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    media_type = (
        "application/vnd.openxmlformats-"
        "officedocument.spreadsheetml.sheet"
    )
    return StreamingResponse(
        buf,
        media_type=media_type,
        headers={
            "Content-Disposition":
                f"attachment; filename=orden_entrada_{o.id}.xlsx"
        },
    )
